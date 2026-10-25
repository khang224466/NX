"""Import the user-supplied 2025-2026 admissions workbook into Nexora's static catalog.

Run with the bundled Python runtime and openpyxl installed. Source cells are preserved
as data; generic source claims and institution homepages are not treated as evidence.
"""

from __future__ import annotations

from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
import json
import re
import sys

from openpyxl import load_workbook


DEFAULT_SOURCE = (
    Path.home()
    / "Downloads"
    / "tong_hop_270_dai_hoc_hoc_vien_viet_nam_2025_2026 (1).xlsx"
)
OFFICIAL_YDS_2026_NOTICE = (
    "https://admin.ump.edu.vn/uploads/ckeditor/files/Truong/DaoTaoDaiHoc/"
    "TuyenSinh/2026/2415_ThongBao_XetTuyenDHCQ2026.pdf"
)
OFFICIAL_VNAM_2026_TUITION = "https://vnam.edu.vn/NewsDetail.aspx?ItemID=2090"


def text_cell(sheet, row: int, column: int) -> str:
    value = sheet.cell(row, column).value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def extract_combinations(value: str) -> list[str]:
    return list(
        dict.fromkeys(
            re.findall(r"\b(?:[A-Z]\d{2}|M\d{2}|T\d{2}|X\d{2})\b", value.upper())
        )
    )


def extract_program_code(value: str) -> str | None:
    found = set(re.findall(r"(?<!\d)(\d{6,7})(?!\d)", value))
    return next(iter(found)) if len(found) == 1 else None


def extract_single_annual_fee(value: str) -> int | None:
    match = re.match(
        r"^\s*(\d{1,3}(?:[.,]\d{3})+|\d+)\s*(?:đồng|vnđ)", value, re.I
    )
    if not match:
        return None
    try:
        return int(re.sub(r"\D", "", match.group(1)))
    except ValueError:
        return None


def valid_http_url(value: str) -> str | None:
    return value if re.match(r"^https?://", value, re.I) else None


def get_official_yds_fees() -> list[dict]:
    """Published 2026 admissions estimates from the university's official notice."""
    rows = [
        ("7720101", "Y khoa", 88_000_000),
        ("7720110", "Y học dự phòng", 55_000_000),
        ("7720115", "Y học cổ truyền", 55_000_000),
        ("7720201", "Dược học", 63_500_000),
        ("7720203", "Hóa dược", 52_500_000),
        ("7720301", "Điều dưỡng", 48_000_000),
        ("7720301_03", "Điều dưỡng chuyên ngành Gây mê hồi sức", 48_000_000),
        ("7720302", "Hộ sinh", 48_000_000),
        ("7720401", "Dinh dưỡng", 48_000_000),
        ("7720501", "Răng - Hàm - Mặt", 90_000_000),
        ("7720502", "Kỹ thuật phục hình răng", 50_000_000),
        ("7720601", "Kỹ thuật xét nghiệm y học", 48_000_000),
        ("7720602", "Kỹ thuật hình ảnh y học", 48_000_000),
        ("7720603", "Kỹ thuật phục hồi chức năng", 48_000_000),
        ("7720701", "Y tế công cộng", 46_000_000),
        ("7760101", "Công tác xã hội", 30_000_000),
        (None, "Công nghệ dược phẩm", 53_500_000),
        ("7310401", "Tâm lý học", 30_000_000),
    ]
    return [
        {
            "institutionCode": "YDS",
            "institutionName": "Đại học Y Dược TP. Hồ Chí Minh",
            "programCode": code,
            "programName": name,
            "academicYear": "2026-2027",
            "annualVnd": amount,
            "feeType": "official_admission_notice_estimate",
            "verificationStatus": "OFFICIAL_PUBLISHED_ESTIMATE",
            "sourceTitle": "Thông báo xét tuyển đại học hệ chính quy năm 2026 (01/06/2026)",
            "sourceUrl": OFFICIAL_YDS_2026_NOTICE,
        }
        for code, name, amount in rows
    ]


def get_official_vnam_fees() -> list[dict]:
    """Published tuition decision rates for VNAM full-time Thanh nhạc."""
    return [
        {
            "institutionCode": "NVH",
            "institutionName": "Học viện Âm nhạc Quốc gia Việt Nam",
            "programCode": "7210205",
            "programName": "Thanh nhạc",
            "academicYear": "2026-2027",
            "annualVnd": 17_100_000,
            "feeType": "official_tuition_decision_rate",
            "verificationStatus": "OFFICIAL_PUBLISHED_RATE",
            "sourceTitle": "Quyết định về mức thu học phí năm học 2026-2027",
            "sourceUrl": OFFICIAL_VNAM_2026_TUITION,
        }
    ]


def import_workbook(source_path: Path, output_path: Path) -> dict:
    workbook = load_workbook(source_path, data_only=True, read_only=False)

    institutions = []
    sheet = workbook["2_Toan_Bo_270_Don_Vi"]
    for row in range(5, sheet.max_row + 1):
        source_id = text_cell(sheet, row, 1)
        if not source_id:
            continue
        institutions.append(
            {
                "sourceId": source_id,
                "ownershipTypeRaw": text_cell(sheet, row, 2),
                "governingBodyRaw": text_cell(sheet, row, 3),
                "name": text_cell(sheet, row, 4),
                "institutionCodeRaw": text_cell(sheet, row, 5),
                "locationRaw": text_cell(sheet, row, 6),
                "focusAreasRaw": text_cell(sheet, row, 7),
                "admissionCombinations": extract_combinations(text_cell(sheet, row, 8)),
                "cutoff2025Raw": text_cell(sheet, row, 9),
                "cutoff2026Raw": text_cell(sheet, row, 10),
                "tuition2026_2027Raw": text_cell(sheet, row, 11),
                "officialHomepage": valid_http_url(text_cell(sheet, row, 12)),
                "sourceSearchQueryRaw": text_cell(sheet, row, 13),
                "sourceClaimRaw": text_cell(sheet, row, 14),
                "sourceSheet": sheet.title,
                "sourceRow": row,
                "verificationStatus": "SUPPLIED_UNVERIFIED",
            }
        )

    programs = []
    for sheet_name, kind in [
        ("6_Chi_Tiet_Nganh_DHQG_HN_HCM", "national_university_detail"),
        ("7_Chi_Tiet_Nganh_CTU_UDN", "ctu_danang_detail"),
    ]:
        sheet = workbook[sheet_name]
        for row in range(5, sheet.max_row + 1):
            source_id = text_cell(sheet, row, 1)
            if not source_id:
                continue
            if kind == "national_university_detail":
                sector, institution, raw_code, name, combinations, quota, cutoff, tuition = [
                    text_cell(sheet, row, column) for column in range(2, 9 + 1)
                ]
                programs.append(
                    {
                        "sourceId": source_id,
                        "programKind": kind,
                        "sectorRaw": sector,
                        "institutionNameRaw": institution,
                        "institutionCodeRaw": None,
                        "rawProgramCode": raw_code,
                        "parsedProgramCode": extract_program_code(raw_code),
                        "programName": name,
                        "admissionCombinations": extract_combinations(combinations),
                        "admissionMethodsRaw": None,
                        "quotaRaw": quota,
                        "cutoffRaw": cutoff,
                        "tuitionRaw": tuition,
                        "sourceReferenceRaw": None,
                        "sourceUrl": None,
                        "sourceSheet": sheet_name,
                        "sourceRow": row,
                        "verificationStatus": "SUPPLIED_UNVERIFIED",
                    }
                )
            else:
                institution, raw_code, name, methods, quota, cutoff, tuition, source_ref = [
                    text_cell(sheet, row, column) for column in range(2, 10)
                ]
                match = re.search(r"\b(?:CTU|UDN|TCT|DND)\b", institution, re.I)
                programs.append(
                    {
                        "sourceId": source_id,
                        "programKind": kind,
                        "institutionNameRaw": institution,
                        "institutionCodeRaw": match.group(0).upper() if match else None,
                        "rawProgramCode": raw_code,
                        "parsedProgramCode": extract_program_code(raw_code),
                        "programName": name,
                        "admissionCombinations": extract_combinations(methods),
                        "admissionMethodsRaw": methods,
                        "quotaRaw": quota,
                        "cutoffRaw": cutoff,
                        "tuitionRaw": tuition,
                        "sourceReferenceRaw": source_ref,
                        "sourceUrl": valid_http_url(source_ref),
                        "sourceSheet": sheet_name,
                        "sourceRow": row,
                        "verificationStatus": "SUPPLIED_UNVERIFIED",
                    }
                )

    official_facts = get_official_yds_fees() + get_official_vnam_fees()
    official_by_code = {
        row["programCode"]: row
        for row in official_facts
        if row["programCode"] is not None
    }
    tuition_records = []
    sheet = workbook["8_Hoc_Phi_Chinh_Xac_Tung_Nganh"]
    for row in range(5, sheet.max_row + 1):
        source_id = text_cell(sheet, row, 1)
        if not source_id:
            continue
        (
            institution_name,
            institution_code,
            raw_code,
            program_name,
            study_track,
            raw_annual,
            raw_monthly,
            source_claim,
        ) = [text_cell(sheet, row, column) for column in range(2, 10)]
        code = extract_program_code(raw_code)
        reported = extract_single_annual_fee(raw_annual)
        is_yds = institution_code.upper() == "YDS"
        official = official_by_code.get(code) if is_yds and code else None

        if is_yds and code == "7310401" and "Công tác xã hội" in program_name:
            status = "AMBIGUOUS_PROGRAM_GROUP"
            note = (
                "Workbook row groups Psychology and Social Work under one program code; "
                "the official notice lists these programs separately. Use the separate "
                "official facts rather than mapping both names to this one code."
            )
            resolved = None
            verified_url = OFFICIAL_YDS_2026_NOTICE
        elif is_yds and code == "7720601":
            status = "AMBIGUOUS_AGGREGATE"
            note = (
                "Workbook groups several programs under one row and fee; the official 2026 "
                "notice lists different fees for some programs in that group."
            )
            resolved = None
            verified_url = OFFICIAL_YDS_2026_NOTICE
        elif official:
            resolved = official["annualVnd"]
            status = (
                "OFFICIAL_SOURCE_MATCH"
                if reported == resolved
                else "OFFICIAL_SOURCE_CORRECTION"
            )
            note = (
                "Matches the official published estimate."
                if reported == resolved
                else f"Workbook value {reported} VND differs; use official published estimate {resolved} VND."
            )
            verified_url = OFFICIAL_YDS_2026_NOTICE
        else:
            resolved = None
            status = "SUPPLIED_UNVERIFIED"
            note = (
                "Workbook assertion has no row-specific source URL; retain as supplied "
                "data pending verification."
            )
            verified_url = None

        tuition_records.append(
            {
                "sourceId": source_id,
                "institutionNameRaw": institution_name,
                "institutionCodeRaw": institution_code,
                "rawProgramCode": raw_code,
                "parsedProgramCode": code,
                "programNameRaw": program_name,
                "studyTrackRaw": study_track,
                "reportedAnnualRaw": raw_annual,
                "reportedAnnualVnd": reported,
                "reportedMonthlyRaw": raw_monthly,
                "sourceClaimRaw": source_claim,
                "sourceSheet": sheet.title,
                "sourceRow": row,
                "verificationStatus": status,
                "resolvedAnnualVnd": resolved,
                "resolutionNote": note,
                "verificationSourceUrl": verified_url,
            }
        )

    source_hash = sha256(source_path.read_bytes()).hexdigest()
    data = {
        "schemaVersion": 1,
        "datasetId": "nexora-admissions-workbook-2025-2026-import-2026-10-23",
        "importedAt": date.today().isoformat(),
        "source": {
            "fileName": source_path.name,
            "sha256": source_hash,
            "workbookLabel": (
                "2025-2026; detailed admissions and tuition sections include 2026 "
                "and 2026-2027 claims"
            ),
            "notes": (
                "User-supplied workbook. Generic source claims and institution "
                "homepages are not treated as row-specific evidence."
            ),
        },
        "reconciliationPolicy": {
            "verifiedOfficialSourcesTakePrecedence": True,
            "unverifiedConflictsAreRetainedButNotPromoted": True,
            "segmentSheets3to5AreNotDuplicated": True,
            "programCutoffsAndTuitionRemainRawUntilVerified": True,
        },
        "institutions": institutions,
        "programs": programs,
        "tuitionRecords": tuition_records,
        "verifiedTuitionFacts": official_facts,
        "dataQuality": {
            "institutionRows": len(institutions),
            "programRows": len(programs),
            "tuitionRows": len(tuition_records),
            "tuitionRowsWithSingleParsedAmount": sum(
                row["reportedAnnualVnd"] is not None for row in tuition_records
            ),
            "tuitionOfficialMatches": sum(
                row["verificationStatus"] == "OFFICIAL_SOURCE_MATCH"
                for row in tuition_records
            ),
            "tuitionOfficialCorrections": sum(
                row["verificationStatus"] == "OFFICIAL_SOURCE_CORRECTION"
                for row in tuition_records
            ),
            "tuitionAmbiguousAggregates": sum(
                row["verificationStatus"] == "AMBIGUOUS_AGGREGATE"
                for row in tuition_records
            ),
            "tuitionAmbiguousProgramGroups": sum(
                row["verificationStatus"] == "AMBIGUOUS_PROGRAM_GROUP"
                for row in tuition_records
            ),
            "sourceCoverage": (
                "Official source-backed tuition includes YDS 2026 admission estimates "
                "and the VNAM 2026-2027 tuition decision for Thanh nhạc. Other supplied "
                "workbook details remain unverified."
            ),
        },
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return data


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parents[1]
    source_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    output_path = project_root / "src" / "data" / "importedAdmissions2025_2026.json"
    result = import_workbook(source_path, output_path)
    print(
        json.dumps(
            {
                "output": str(output_path),
                "bytes": output_path.stat().st_size,
                "institutions": len(result["institutions"]),
                "programs": len(result["programs"]),
                "tuitionRows": len(result["tuitionRecords"]),
                "parsedTuition": result["dataQuality"]["tuitionRowsWithSingleParsedAmount"],
                "officialMatches": result["dataQuality"]["tuitionOfficialMatches"],
                "officialCorrections": result["dataQuality"]["tuitionOfficialCorrections"],
                "ambiguousAggregates": result["dataQuality"]["tuitionAmbiguousAggregates"],
                "ambiguousProgramGroups": result["dataQuality"]["tuitionAmbiguousProgramGroups"],
            },
            ensure_ascii=False,
        )
    )
